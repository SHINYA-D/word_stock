#!/usr/bin/env python3
"""テスト項目書（MD マトリクス）→ Excel 変換。

入力:
  test/test_cases/**/*.md          … 各テストエージェントが生成した項目書
  coverage/harness_report.json     … 最新のハーネス結果（任意。あればサマリ・要確認一覧に反映）
  .test_loop/state.json            … test-loop の進捗（任意。production_bugs を要確認一覧に反映）
                                     旧 coverage/test_loop_state.json も後方互換で読む

出力:
  ~/Desktop/WordStock_テスト項目書_YYYYMMDD.xlsx   （--out で変更可）

シート構成（左から）:
  表紙            … タイトル・基本情報・実施状況サマリ（自動集計数式）
  記載要領・観点一覧 … 各列の意味・観点分類（正常系/異常系/境界値）の考え方
  サマリ          … 対象ファイルごとのテスト数・カテゴリ内訳・カバレッジ
  要確認一覧       … テスト工程で解消しきれなかった問題（理由なし未達 / 90%未満（理由あり） /
                     テスト失敗 / テスト漏れ / プロダクションコードのバグ）
  単体テスト項目書  … test()  ベースの項目一覧（1テストケース=1行）
  Widgetテスト項目書 … testWidgets() ベースの項目一覧（1テストケース=1行）
  対象外一覧       … 登録簿(test/coverage_exclusions.txt) + 項目書MD の対象外

出力フォーマットは `/Users/a12345/Desktop/chousa/generate_test_item_template.py` の
表紙／記載要領・観点一覧／テスト項目書シートのレイアウトに統一している
（スタイル定義・シート構成をこのスクリプトへ移植）。

MD フォーマット規約は .claude/skills/excel-testdoc-authoring/SKILL.md を参照。
既存の test/test_cases/infrastructure/repositories/folder_repository_impl_test_cases.md
を正とする。
"""

from __future__ import annotations

import argparse
import glob
import json
import os
import re
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    raise SystemExit(
        "openpyxl が見つかりません。\n"
        "  python3 -m venv scripts/.venv\n"
        "  scripts/.venv/bin/pip install -r scripts/requirements.txt\n"
        "を実行してから scripts/.venv/bin/python3 で起動してください。"
    )

import loop_state

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
TEST_CASES_DIR = os.path.join(REPO, "test", "test_cases")

# --------------------------------------------------------------------------- #
# 共通スタイル定義（generate_test_item_template.py から移植）
# --------------------------------------------------------------------------- #
FONT_NAME = "游ゴシック"

TITLE_FONT = Font(name=FONT_NAME, size=16, bold=True)
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, color="FF666666")
SECTION_FONT = Font(name=FONT_NAME, size=10, bold=True)
SECTION_FONT_L = Font(name=FONT_NAME, size=13, bold=True)
LABEL_FONT = Font(name=FONT_NAME, size=10)
NOTE_FONT = Font(name=FONT_NAME, size=9, italic=True, color="FF666666")
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=10)

LABEL_FILL = PatternFill("solid", fgColor="FFD9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
HEADER_FILL = PatternFill("solid", fgColor="FF305496")
CATEGORY_FILL = {
    "正常系": PatternFill("solid", fgColor="E3F2E1"),
    "異常系": PatternFill("solid", fgColor="FBE3E0"),
    "エッジケース": PatternFill("solid", fgColor="FFF3D6"),
    "境界値": PatternFill("solid", fgColor="FFF3D6"),
}

RED_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFC00000")
RED_FILL = PatternFill("solid", fgColor="FFFBE3E0")
YELLOW_FILL = PatternFill("solid", fgColor="FFFFF3D6")

# サブエージェントが内部リトライ上限で打ち切った未カバー行に付ける目印。
# 理由が確定していないので「理由あり」とは扱わない（定義は loop_state.py と共有）。
PENDING_MARK = loop_state.PENDING_MARK

THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center")
WRAP_CENTER_H = Alignment(wrap_text=True, vertical="center", horizontal="center")

# 「単体テスト項目書」「Widgetテスト項目書」シートの1件目のデータ行
ITEM_DATA_FIRST_ROW = 6


def style_cell(cell, font=BODY_FONT, fill=None, align=WRAP_TOP, border=BORDER):
    cell.font = font
    if fill is not None:
        cell.fill = fill
    cell.alignment = align
    cell.border = border


# --------------------------------------------------------------------------- #
# MD パース
# --------------------------------------------------------------------------- #
def read_md_tables(text: str):
    """MD 内の全パイプ表を [(header:list[str], rows:list[list[str]])] で返す。"""
    tables = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\|[\s:|-]+\|?$", lines[i + 1].strip()
        ):
            header = [c.strip() for c in line.strip("|").split("|")]
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            tables.append((header, rows))
            i = j
        else:
            i += 1
    return tables


def meta_from_tables(tables):
    """「| 項目 | 値 |」形式の縦持ち表から辞書を作る。"""
    meta = {}
    for header, rows in tables:
        if len(header) == 2 and header[0] in ("項目",):
            for r in rows:
                if len(r) == 2:
                    meta[r[0]] = r[1]
    return meta


def find_case_table(tables):
    """「テストケース一覧」表（テスト名 / カテゴリ 等を含む）を探す。"""
    for header, rows in tables:
        joined = " ".join(header)
        if "テスト名" in joined:
            return header, rows
    return None, None


# 「対象外」を含む見出し配下の箇条書きを (項目, 理由) で返す。
# ループ側（loop_state.compute_verdict）と同じ解釈にするため実装を共有する。
parse_offtargets = loop_state.parse_offtargets


DETAIL_FIELD_ALIASES = {
    "prereq": ["事前条件"],
    "input": ["入力値・テスト条件", "入力条件"],
    "steps": ["操作手順"],
    "expected": ["期待結果", "期待値"],
}


def parse_detail_sections(text: str):
    """「## テストケース詳細」配下の各ケースブロックから
    事前条件/入力値・テスト条件/操作手順/期待結果 を抽出し、
    ケース番号(文字列)をキーにした dict を返す。
    旧フィールド名（入力条件/期待値）にもフォールバックする。
    """
    m = re.search(r"^##\s*テストケース詳細\s*$", text, re.M)
    if not m:
        return {}
    tail = text[m.end():]
    m2 = re.search(r"^##\s+\S", tail, re.M)
    if m2:
        tail = tail[:m2.start()]

    case_matches = list(re.finditer(r"^###\s*テストケース(\d+)[:：]?", tail, re.M))
    results = {}
    for i, cm in enumerate(case_matches):
        start = cm.end()
        end = case_matches[i + 1].start() if i + 1 < len(case_matches) else len(tail)
        block = tail[start:end]
        raw_fields = {}
        for line in block.splitlines():
            bm = re.match(r"^-\s*\*\*(.+?)\*\*[:：]\s*(.*)$", line.strip())
            if bm:
                raw_fields[bm.group(1).strip()] = bm.group(2).strip()

        fields = {}
        for key, aliases in DETAIL_FIELD_ALIASES.items():
            for alias in aliases:
                if alias in raw_fields and raw_fields[alias]:
                    fields[key] = raw_fields[alias]
                    break
            else:
                fields[key] = ""
        results[cm.group(1)] = fields
    return results


def classify(path: str) -> str:
    """単体 / Widget の振り分け。"""
    name = os.path.basename(path)
    if name.endswith("_page_test_cases.md"):
        return "Widget"
    if "testWidgets" in name:
        return "Widget"
    return "単体"


def parse_file(path: str):
    with open(path, encoding="utf-8") as f:
        text = f.read()
    tables = read_md_tables(text)
    meta = meta_from_tables(tables)
    header, rows = find_case_table(tables)
    target = (
        meta.get("ファイルパス")
        or meta.get("対象ファイル")
        or os.path.relpath(path, TEST_CASES_DIR)
    )
    cls = meta.get("クラス名") or meta.get("ViewModelクラス") or meta.get("ページ名") or ""
    details = parse_detail_sections(text)

    def col(hrow, *names):
        for n in names:
            if n in hrow:
                return hrow.index(n)
        return None

    cases = []
    if header and rows:
        i_no = col(header, "#", "No", "No.")
        i_name = col(header, "テスト名")
        i_cat = col(header, "カテゴリ", "状態パターン")
        i_method = col(header, "対象メソッド", "検証項目")
        i_status = col(header, "状態")
        for idx, r in enumerate(rows, start=1):
            def get(hidx):
                return r[hidx] if hidx is not None and hidx < len(r) else ""
            case_no = get(i_no) or str(idx)
            detail = details.get(case_no, {})
            cases.append({
                "no": case_no,
                "name": get(i_name),
                "category": get(i_cat) or "-",
                "method": get(i_method),
                "status": get(i_status) or "✅",
                "prereq": detail.get("prereq", ""),
                "input": detail.get("input", ""),
                "steps": detail.get("steps", ""),
                "expected": detail.get("expected", ""),
            })
    return {
        "doc": os.path.relpath(path, REPO),
        "kind": classify(path),
        "target": target,
        "class": cls,
        "cases": cases,
        "offtargets": parse_offtargets(text),
    }


# --------------------------------------------------------------------------- #
# シート1：表紙
# --------------------------------------------------------------------------- #
def build_cover_sheet(wb, generated_at: str, total_cases: int):
    ws = wb.active
    ws.title = "表紙"
    ws.sheet_view.showGridLines = False

    for col, w in zip("ABCDE", [4, 22, 40, 4, 22]):
        ws.column_dimensions[col].width = w

    ws.row_dimensions[2].height = 19.7
    for r in list(range(3, 4)) + list(range(5, 14)):
        ws.row_dimensions[r].height = 15

    ws["B2"] = "WordStock テスト項目書"
    ws["B2"].font = TITLE_FONT
    ws.merge_cells("B2:F2")

    ws["B3"] = "（Unit / Widget Test Specification）"
    ws["B3"].font = SUBTITLE_FONT
    ws.merge_cells("B3:F3")

    ws["B5"] = "■ 基本情報"
    ws["B5"].font = SECTION_FONT
    ws.merge_cells("B5:C5")
    ws["E5"] = "■ 実施状況サマリ（自動集計）"
    ws["E5"].font = SECTION_FONT
    ws.merge_cells("E5:F5")

    basic_info = [
        ("プロジェクト名", "WordStock"),
        ("対象システム名", "WordStock"),
        ("対象モジュール／機能", ""),
        ("版数", "1.0"),
        ("作成者", ""),
        ("作成日", generated_at),
        ("承認者", ""),
        ("承認日", ""),
    ]
    row = 6
    for label, value in basic_info:
        style_cell(ws.cell(row=row, column=2, value=label), font=LABEL_FONT,
                   fill=LABEL_FILL, align=WRAP_CENTER)
        style_cell(ws.cell(row=row, column=3, value=value or None), font=BODY_FONT,
                   fill=INPUT_FILL, align=WRAP_CENTER)
        row += 1

    # 実施状況サマリ（単体テスト項目書 + Widgetテスト項目書 の両シートを合算）
    s = ITEM_DATA_FIRST_ROW
    unit = "単体テスト項目書"
    widget = "Widgetテスト項目書"
    total_formula = (
        f"=COUNTA({unit}!B{s}:B100000)+COUNTA({widget}!B{s}:B100000)"
    )
    done_formula = (
        f'=COUNTIF({unit}!M{s}:M100000,"OK")+COUNTIF({unit}!M{s}:M100000,"NG")'
        f'+COUNTIF({widget}!M{s}:M100000,"OK")+COUNTIF({widget}!M{s}:M100000,"NG")'
    )
    ok_formula = (
        f'=COUNTIF({unit}!M{s}:M100000,"OK")+COUNTIF({widget}!M{s}:M100000,"OK")'
    )
    ng_formula = (
        f'=COUNTIF({unit}!M{s}:M100000,"NG")+COUNTIF({widget}!M{s}:M100000,"NG")'
    )
    pending_formula = (
        f'=COUNTIF({unit}!M{s}:M100000,"未実施")+COUNTIF({widget}!M{s}:M100000,"未実施")'
    )
    rate_formula = (
        f"=IFERROR(({done_formula.lstrip('=')})/({total_formula.lstrip('=')}),0)"
    )
    summary = [
        ("テスト項目総数", total_formula, "General"),
        ("実施済み件数", done_formula, "General"),
        ("OK件数", ok_formula, "General"),
        ("NG件数", ng_formula, "General"),
        ("未実施件数", pending_formula, "General"),
        ("消化率", rate_formula, "0.0%"),
    ]
    row = 6
    for label, formula, numfmt in summary:
        style_cell(ws.cell(row=row, column=5, value=label), font=LABEL_FONT,
                   fill=LABEL_FILL, align=WRAP_CENTER)
        fcell = ws.cell(row=row, column=6, value=formula)
        style_cell(fcell, font=BODY_FONT, fill=None, align=WRAP_CENTER_H)
        fcell.number_format = numfmt
        row += 1

    usage_note = (
        "【本書の使い方】\n"
        "・本書は test/test_cases/**/*.md（各テストエージェントが生成した項目書）を"
        "自動集計して生成しています。\n"
        "・「単体テスト項目書」「Widgetテスト項目書」シートに実際のテストケースが1行ずつ並びます"
        "（黄色セルは手入力欄の目安ではなく、記入例の色分けは行っていません）。\n"
        "・「記載要領・観点一覧」シートに、各列の意味とテスト観点（正常系／異常系／境界値）の考え方をまとめています。\n"
        "・「結果」列は各MDの「状態」列（✅=ハーネスでテスト到達・実装済み）から自動的にOK/未実施を設定しています。"
        "手動で再判定した場合はOK／NG／未実施に書き換えてください。表紙のサマリは自動集計されます"
        "（値を変更後、Excelで再計算してください）。\n"
        f"・生成日時: {generated_at}　総テストケース数: {total_cases}\n"
        "・本ファイルは scripts/gen_test_excel.py から再生成できます。"
    )
    ws["B14"] = usage_note
    ws["B14"].font = LABEL_FONT
    ws["B14"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B14:F20")


# --------------------------------------------------------------------------- #
# シート2：記載要領・観点一覧
# --------------------------------------------------------------------------- #
def build_guideline_sheet(wb):
    ws = wb.create_sheet("記載要領・観点一覧")
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABC", [4, 18, 60]):
        ws.column_dimensions[col].width = w

    ws.row_dimensions[2].height = 16.15

    ws["B2"] = "記載要領（各列の意味）"
    ws["B2"].font = SECTION_FONT_L
    ws.merge_cells("B2:C2")

    ws.row_dimensions[4].height = 17.15
    style_cell(ws["B4"], font=HEADER_FONT, fill=HEADER_FILL, align=WRAP_CENTER_H)
    ws["B4"] = "列名"
    style_cell(ws["C4"], font=HEADER_FONT, fill=HEADER_FILL, align=WRAP_CENTER_H)
    ws["C4"] = "記載内容"

    col_defs = [
        ("No.", "テスト項目の通し番号（シート内の通し番号）。"),
        ("テストID", "テストケースを一意に識別するID（例：UT-画面名-001／Widgetは WT-）。管理・追跡に使用する。"),
        ("大分類（対象機能／クラス）", "テスト対象の機能名・画面名・クラス名など、大きな単位を記載する。"),
        ("中分類（対象メソッド／画面項目）", "テスト対象のメソッド名・関数名・画面項目など、より具体的な単位を記載する。"),
        ("観点分類", "テストの種類を「正常系／異常系／境界値」から選択する（詳細は下表）。"),
        ("テスト観点（確認内容）", "そのテストで何を確認したいのかを、一読して目的が分かるように具体的に記載する。"),
        ("事前条件", "テスト実施前に満たしているべき状態（Fakeの初期状態、ログイン済み 等）。"),
        ("入力値・テスト条件", "実際に与える入力データやパラメータ、操作条件を具体的に記載する。"),
        ("操作手順", "期待結果が得られるまでの操作・実行手順を、誰が読んでも再現できるように記載する。"),
        ("期待結果", "「正しく処理されること」等の曖昧な表現は避け、戻り値・画面表示・DBの状態など判定可能な形で具体的に記載する。"),
        ("実施日／実施者", "テストを実施した日付と担当者名を記録する。"),
        ("結果", "実施結果（OK／NG／未実施）。MDの「状態」列（✅=ハーネスでテスト到達・実装済み）から自動設定される。"
                "手動で再判定した場合は書き換えてよい。"),
        ("備考", "NG時の不具合番号、補足事項、再テスト結果などを記載する。"),
    ]
    row = 5
    for name, desc in col_defs:
        style_cell(ws.cell(row=row, column=2, value=name), font=BODY_FONT, align=WRAP_TOP)
        style_cell(ws.cell(row=row, column=3, value=desc), font=BODY_FONT, align=WRAP_TOP)
        ws.row_dimensions[row].height = 30
        row += 1

    row += 2
    section_row = row
    ws.cell(row=section_row, column=2, value="テスト観点分類の考え方").font = SECTION_FONT_L
    ws.merge_cells(start_row=section_row, start_column=2, end_row=section_row, end_column=3)
    ws.row_dimensions[section_row].height = 16.15

    header_row = section_row + 2
    ws.row_dimensions[header_row].height = 17.15
    style_cell(ws.cell(row=header_row, column=2, value="観点分類"), font=HEADER_FONT,
               fill=HEADER_FILL, align=WRAP_CENTER_H)
    style_cell(ws.cell(row=header_row, column=3, value="内容・考え方"), font=HEADER_FONT,
               fill=HEADER_FILL, align=WRAP_CENTER_H)

    perspectives = [
        ("正常系", "仕様通りの入力・操作を行った場合に、期待通りの結果が得られることを確認する（代表値による同値分割を含む）。"),
        ("異常系", "仕様上想定される誤った入力・操作（未入力、不正な形式、例外発生 等）に対し、"
                 "適切なエラー処理（Failure変換等）が行われることを確認する。"),
        ("境界値", "入力値の有効範囲の境界（最小値／最大値／その前後／ネストの深さ 等）を対象に、"
                 "正しく処理される／エラーになることを確認する（境界値分析。エッジケースを含む）。"),
    ]
    row = header_row + 1
    for name, desc in perspectives:
        style_cell(ws.cell(row=row, column=2, value=name), font=BODY_FONT, align=WRAP_CENTER_H)
        style_cell(ws.cell(row=row, column=3, value=desc), font=BODY_FONT, align=WRAP_TOP)
        ws.row_dimensions[row].height = 39.75
        row += 1

    row += 2
    note_row = row
    note = ws.cell(
        row=note_row, column=2,
        value=(
            "本シートは各テストエージェントが生成する test/test_cases/**/*.md の記載要領と対応している。"
            "MDフォーマットの詳細規約は .claude/skills/excel-testdoc-authoring/SKILL.md を参照。"
        ),
    )
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row + 3, end_column=3)


# --------------------------------------------------------------------------- #
# シート3・4：単体テスト項目書 / Widgetテスト項目書（1テストケース=1行）
# --------------------------------------------------------------------------- #
def make_slug(fi) -> str:
    name = fi["class"] or os.path.basename(fi["target"]).replace(".dart", "")
    slug = re.sub(r"[^A-Za-z0-9]", "", name).upper()
    return slug or "TARGET"


def map_status_to_result(status: str) -> str:
    """MD の「状態」列（✅=ハーネスでテスト到達・実装済み）を結果列(OK/NG/未実施)へ変換する。"""
    s = (status or "").strip()
    if s in ("✅", "OK", "ok"):
        return "OK"
    if s in ("❌", "NG", "ng"):
        return "NG"
    return "未実施"


def build_test_item_sheet(ws, files, id_prefix: str, sheet_title: str, generated_at: str):
    ws.sheet_view.showGridLines = False

    col_widths = {
        "A": 5, "B": 16, "C": 18, "D": 20, "E": 10, "F": 26, "G": 22,
        "H": 24, "I": 24, "J": 26, "K": 11, "L": 10, "M": 8, "N": 18,
    }
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    ws.row_dimensions[1].height = 24
    ws["A1"] = sheet_title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:N1")

    total_cases = sum(len(fi["cases"]) for fi in files)
    ws["A2"] = f"対象ファイル数：{len(files)}　　テストケース数：{total_cases}"
    ws["A2"].font = NOTE_FONT
    ws.merge_cells("A2:N2")

    headers = [
        "No.", "テストID", "大分類\n（対象機能／クラス）", "中分類\n（対象メソッド／画面項目）",
        "観点分類", "テスト観点（確認内容）", "事前条件", "入力値・テスト条件", "操作手順",
        "期待結果", "実施日", "実施者", "結果", "備考",
    ]
    header_row = 4
    for i, h in enumerate(headers, start=1):
        style_cell(ws.cell(row=header_row, column=i, value=h), font=HEADER_FONT,
                   fill=HEADER_FILL, align=WRAP_CENTER_H)

    note_row = 5
    ws.cell(
        row=note_row, column=1,
        value="観点分類：正常系／異常系／境界値　　結果：OK／NG／未実施　　"
              "（詳しい定義は「記載要領・観点一覧」シートを参照）",
    )
    ws.cell(row=note_row, column=1).font = NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=14)

    center_cols = {1, 5, 11, 12, 13}
    r = ITEM_DATA_FIRST_ROW
    no = 0
    for fi in files:
        slug = make_slug(fi)
        for case in fi["cases"]:
            no += 1
            case_no = case["no"]
            id_no = case_no.zfill(3) if case_no.isdigit() else case_no
            test_id = f"{id_prefix}-{slug}-{id_no}"
            result = map_status_to_result(case["status"])
            jissha = "自動(harness)" if result != "未実施" else ""
            jisshibi = generated_at if result != "未実施" else ""
            row_vals = [
                no, test_id, fi["class"] or fi["target"], case["method"],
                case["category"], case["name"], case["prereq"], case["input"],
                case["steps"], case["expected"], jisshibi, jissha, result, "",
            ]
            for i, val in enumerate(row_vals, start=1):
                align = WRAP_CENTER_H if i in center_cols else WRAP_TOP
                style_cell(ws.cell(row=r, column=i, value=val if val not in ("", None) else None),
                           font=BODY_FONT, align=align)
            fill = CATEGORY_FILL.get(case["category"])
            if fill:
                ws.cell(row=r, column=5).fill = fill
            ws.row_dimensions[r].height = 30
            r += 1

    last_row = r - 1
    if last_row >= ITEM_DATA_FIRST_ROW:
        dv_kanten = DataValidation(type="list", formula1='"正常系,異常系,境界値"', allow_blank=True)
        ws.add_data_validation(dv_kanten)
        dv_kanten.add(f"E{ITEM_DATA_FIRST_ROW}:E{last_row}")

        dv_result = DataValidation(type="list", formula1='"OK,NG,未実施"', allow_blank=True)
        ws.add_data_validation(dv_result)
        dv_result.add(f"M{ITEM_DATA_FIRST_ROW}:M{last_row}")

    ws.freeze_panes = f"A{ITEM_DATA_FIRST_ROW}"


# --------------------------------------------------------------------------- #
# シート：サマリ
# --------------------------------------------------------------------------- #
def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _norm(p: str) -> str:
    p = (p or "").replace("\\", "/")
    i = p.find("lib/")
    return p[i:] if i != -1 else p


def build_summary_sheet(ws, files, report, warnings):
    cov = (report or {}).get("coverage", {})
    tests = (report or {}).get("tests", {})
    threshold = cov.get("threshold", 90.0)
    ws.append(["WordStock テスト項目書", ""])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append(["生成日時", datetime.now().strftime("%Y-%m-%d %H:%M")])
    if report:
        ws.append(["ハーネス結果",
                   f"{tests.get('passed', 0)} passed / {tests.get('failed', 0)} failed "
                   f"/ {tests.get('skipped', 0)} skipped"])
        ws.append(["限定分母カバレッジ",
                   f"{cov.get('overall_pct', '-')}%  "
                   f"({cov.get('lines_hit', '-')}/{cov.get('lines_found', '-')} 行, "
                   f"閾値 {threshold}%)"])
    counts = {}
    for w in warnings:
        counts[w["kind"]] = counts.get(w["kind"], 0) + 1
    ws.append(["要確認",
               " / ".join(f"{k} {counts.get(k, 0)}件" for k, _ in WARNING_KINDS)
               + "　（詳細は「要確認一覧」シート）"])
    if any(counts.get(k, 0) for k, red in WARNING_KINDS if red):
        ws.cell(row=ws.max_row, column=2).font = RED_FONT
    ws.append([])

    hdr_row = ws.max_row + 1
    cols = ["対象ファイル", "種別", "テスト総数", "正常系", "異常系",
            "境界/エッジ", "その他", "カバレッジ", "項目書"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=hdr_row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)

    cov_files = {_norm(f["path"]): f for f in cov.get("files", [])}
    for fi in files:
        cats = [c["category"] for c in fi["cases"]]
        normal = sum(1 for c in cats if "正常" in c or c == "Success")
        abnormal = sum(1 for c in cats if "異常" in c or c == "Error")
        edge = sum(1 for c in cats if "境界" in c or "エッジ" in c)
        other = len(cats) - normal - abnormal - edge
        cf = cov_files.get(_norm(fi["target"]))
        ws.append([
            fi["target"], fi["kind"], len(cats), normal, abnormal, edge, other,
            f"{cf['pct']}%" if cf else "-", fi["doc"],
        ])
        if cf and cf["pct"] < threshold:
            ws.cell(row=ws.max_row, column=8).font = RED_FONT
    autosize(ws, [46, 10, 11, 8, 8, 12, 8, 11, 46])


# --------------------------------------------------------------------------- #
# シート：要確認一覧
# --------------------------------------------------------------------------- #
# (区分, 赤字で強調するか)。サマリの件数表示と並び順もこの順。
WARNING_KINDS = [
    ("理由なし未達", True),
    ("プロダクションコードのバグ", True),
    ("テスト失敗", True),
    ("テスト漏れ", True),
    ("90%未満（理由あり）", False),
    ("対象外に行番号なし", False),
]


def _stem(path: str) -> str:
    """`lib/a/b.dart` / `test/a/b_test.dart` → `b`（テストと対象の突き合わせ用）。"""
    name = os.path.basename(path or "")
    name = name[:-len(".dart")] if name.endswith(".dart") else name
    return name[:-len("_test")] if name.endswith("_test") else name


def collect_warnings(files, report, state, scoped: bool):
    """テスト工程で解消しきれなかった問題を要確認一覧の行として集める。

    scoped=True（--only 指定 or パス指定のハーネス結果）のときは、
    今回の対象ファイル（= files の target）に関係するものだけを載せる。
    """
    cov = (report or {}).get("coverage", {})
    threshold = cov.get("threshold", 90.0)
    targets = {_norm(fi["target"]) for fi in files}
    target_stems = {_stem(t) for t in targets}
    md_by_target = {_norm(fi["target"]): fi for fi in files}
    registered = {_norm(e.get("path", "")) for e in cov.get("exclusions", [])}
    out = []

    # 90% 未満のファイル（理由の有無で区分を分ける）
    for cf in cov.get("files", []):
        path = _norm(cf["path"])
        if cf["pct"] >= threshold or path in registered:
            continue
        if scoped and path not in targets:
            continue
        fi = md_by_target.get(path)
        offtargets = fi["offtargets"] if fi else []
        reasons = [(i, r) for i, r in offtargets if r and PENDING_MARK not in r]
        pending = [(i, r) for i, r in offtargets if PENDING_MARK in r]
        uncovered = ", ".join(map(str, cf.get("uncovered_lines", [])))
        head = f"カバレッジ {cf['pct']}%（未カバー行: {uncovered}）"
        if reasons and not pending:
            out.append({
                "kind": "90%未満（理由あり）", "target": path, "content": head,
                "detail": "\n".join(f"・{i}：{r}" for i, r in reasons),
                "action": "理由が妥当か確認する（対象外一覧シートにも掲載）",
            })
            # 行番号が無いと「その理由が未カバー行を実際に説明しているか」を
            # 機械照合できない（ループ側の verdict も照合を省いて stop している）。
            if not any(loop_state.extract_lines(i) for i, _ in reasons):
                out.append({
                    "kind": "対象外に行番号なし", "target": path,
                    "content": "「## 対象外」の項目に行番号が書かれていない",
                    "detail": "未カバー行: " + uncovered,
                    "action": "項目を「L142-145：理由」の形式にして、"
                              "未カバー行と対応付ける",
                })
        else:
            if not fi:
                detail = "項目書 MD が無い（テストから間接的に読み込まれただけの可能性）"
            elif pending:
                detail = "判断保留のまま残っている行:\n" + "\n".join(
                    f"・{i}" for i, _ in pending)
            else:
                detail = "項目書の「## 対象外」に未カバー行の理由が書かれていない"
            out.append({
                "kind": "理由なし未達", "target": path,
                "content": head + "　⚠ この未カバー行はテストできていない",
                "detail": detail,
                "action": "テストを追加するか、項目書の「## 対象外」に理由を記録する",
            })

    # テスト失敗
    for fl in (report or {}).get("tests", {}).get("failures", []):
        test_file = fl.get("file", "")
        if scoped and test_file and _stem(test_file) not in target_stems:
            continue
        out.append({
            "kind": "テスト失敗", "target": test_file or "-",
            "content": fl.get("name", ""), "detail": fl.get("message", ""),
            "action": "テストコードの誤りか、プロダクションコードのバグかを切り分ける",
        })

    # テスト漏れ（テストも対象外登録も無い対象ファイル）
    for p in cov.get("untested_files", []):
        if scoped and _norm(p) not in targets:
            continue
        out.append({
            "kind": "テスト漏れ", "target": p,
            "content": "テストも対象外登録も無い", "detail": "",
            "action": "テストを書くか test/coverage_exclusions.txt に理由付きで登録する",
        })

    # プロダクションコードのバグ（test-loop が test_loop_state.json に記録したもの）
    for bug in (state or {}).get("production_bugs", []):
        path = _norm(bug.get("path", ""))
        if scoped and path not in targets:
            continue
        line = bug.get("line")
        out.append({
            "kind": "プロダクションコードのバグ",
            "target": f"{path}:{line}" if line else path,
            "content": bug.get("symptom", ""),
            "detail": bug.get("evidence", ""),
            "action": bug.get("recommendation", "") or "プロダクションコードを修正する（テスト工程では変更していない）",
        })

    order = {k: i for i, (k, _) in enumerate(WARNING_KINDS)}
    out.sort(key=lambda w: order[w["kind"]])
    return out


def build_warning_sheet(ws, warnings):
    red_kinds = {k for k, red in WARNING_KINDS if red}
    cols = ["区分", "対象", "内容", "詳細・理由", "対応の目安"]
    ws.append(cols)
    style_header(ws, len(cols))
    if not warnings:
        ws.append(["（なし）", "", "テスト工程で解消しきれなかった問題はありません", "", ""])
    for w in warnings:
        ws.append([w["kind"], w["target"], w["content"], w["detail"], w["action"]])
        red = w["kind"] in red_kinds
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=ws.max_row, column=c)
            cell.alignment = WRAP_TOP
            cell.fill = RED_FILL if red else YELLOW_FILL
            if red and c in (1, 3):
                cell.font = RED_FONT
    autosize(ws, [22, 46, 50, 60, 44])


def build_offtarget_sheet(ws, files, report):
    cols = ["区分", "対象ファイル", "対象外にした項目", "理由"]
    ws.append(cols)
    style_header(ws, len(cols))
    for ex in (report or {}).get("coverage", {}).get("exclusions", []):
        ws.append(["登録簿", ex.get("path", ""), "ファイル全体", ex.get("reason", "")])
        for c in (1, 2, 3, 4):
            ws.cell(row=ws.max_row, column=c).alignment = WRAP_TOP
    for fi in files:
        for item, reason in fi["offtargets"]:
            ws.append(["項目書", fi["target"], item, reason])
            for c in (1, 2, 3, 4):
                ws.cell(row=ws.max_row, column=c).alignment = WRAP_TOP
    autosize(ws, [10, 46, 34, 60])


def main() -> int:
    ap = argparse.ArgumentParser()
    default_out = os.path.join(
        os.path.expanduser("~/Desktop"),
        f"WordStock_テスト項目書_{datetime.now():%Y%m%d}.xlsx",
    )
    ap.add_argument("--out", default=default_out)
    ap.add_argument("--cases-dir", default=TEST_CASES_DIR)
    ap.add_argument("--report", default=os.path.join(REPO, "coverage",
                                                     "harness_report.json"))
    ap.add_argument("--state", default=loop_state.STATE_PATH)
    ap.add_argument(
        "--only", nargs="+", default=None,
        help="この Excel に載せる対象ファイルを限定する（例: word_list_view_model.dart "
             "または lib/presentation/word/word_list_view_model.dart）。"
             "指定した文字列が対象ファイルパス or 項目書パスに部分一致する項目書のみ集約する。"
             "未指定なら test/test_cases/**/*.md を全件集約する（従来動作）。",
    )
    args = ap.parse_args()

    md_paths = sorted(glob.glob(os.path.join(args.cases_dir, "**", "*.md"),
                                recursive=True))
    if not md_paths:
        raise SystemExit(f"項目書 MD が見つかりません: {args.cases_dir}")

    files = [parse_file(p) for p in md_paths]

    if args.only:
        def matches(fi):
            hay = f"{fi['target']} {fi['doc']}".replace("\\", "/")
            return any(needle in hay for needle in args.only)

        filtered = [fi for fi in files if matches(fi)]
        if not filtered:
            raise SystemExit(
                f"--only に一致する項目書が見つかりません: {args.only}\n"
                f"（走査対象: {[f['target'] for f in files]}）"
            )
        files = filtered
    report = None
    if os.path.exists(args.report):
        with open(args.report, encoding="utf-8") as f:
            report = json.load(f)
    state = None
    state_path = args.state
    if not os.path.exists(state_path) and os.path.exists(loop_state.LEGACY_STATE_PATH):
        state_path = loop_state.LEGACY_STATE_PATH  # 旧 coverage/ 配置との後方互換
    if os.path.exists(state_path):
        with open(state_path, encoding="utf-8") as f:
            state = json.load(f)

    # --only 指定、またはハーネスがパス指定で実行されていれば、今回の対象分だけを要確認に載せる
    scoped = bool(args.only) or bool((report or {}).get("paths"))
    warnings = collect_warnings(files, report, state, scoped)

    generated_at = datetime.now().strftime("%Y/%m/%d")
    total_cases = sum(len(f["cases"]) for f in files)

    wb = Workbook()
    build_cover_sheet(wb, generated_at, total_cases)
    build_guideline_sheet(wb)
    build_summary_sheet(wb.create_sheet("サマリ"), files, report, warnings)
    build_warning_sheet(wb.create_sheet("要確認一覧"), warnings)
    build_test_item_sheet(wb.create_sheet("単体テスト項目書"),
                          [f for f in files if f["kind"] == "単体"], "UT", "単体テスト項目書",
                          generated_at)
    build_test_item_sheet(wb.create_sheet("Widgetテスト項目書"),
                          [f for f in files if f["kind"] == "Widget"], "WT", "Widgetテスト項目書",
                          generated_at)
    build_offtarget_sheet(wb.create_sheet("対象外一覧"), files, report)

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    wb.save(args.out)
    print(f"✅ {args.out}")
    print(f"   項目書 {len(files)} 件 / テストケース {total_cases} 件")
    red_kinds = {k for k, red in WARNING_KINDS if red}
    print(f"   要確認 {len(warnings)} 件"
          f"（うち赤字 {sum(1 for w in warnings if w['kind'] in red_kinds)} 件）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
